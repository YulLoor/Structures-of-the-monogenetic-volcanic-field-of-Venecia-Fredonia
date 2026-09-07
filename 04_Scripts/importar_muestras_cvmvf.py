# -*- coding: utf-8 -*-
"""
Importa 01_Insumos/Muestras CVMVF.xlsx a la GDB como puntos
(WGS84 lat/long -> EPSG:9377) y los anade en rojo a Map_Estructural
(y a Mapa_Litologico si existe).

Ejecutar:
  "C:\\Program Files\\ArcGIS\\Pro\\bin\\Python\\envs\\arcgispro-py3\\python.exe" ^
    04_Scripts\\importar_muestras_cvmvf.py
"""

from __future__ import annotations

import os
import sys
import traceback

import arcpy

ROOT = r"C:\PROYECTO_GIS_VF_Antigraviti"
APRX_PATH = os.path.join(ROOT, "Venecia_Fredonia_Analisis_Estructural.aprx")
GDB = os.path.join(
    ROOT,
    "Venecia_Fredonia_Analisis_Estructural",
    "Venecia_Fredonia_Analisis_Estructural.gdb",
)
EXCEL = os.path.join(ROOT, "01_Insumos", "Muestras CVMVF.xlsx")
FC_NAME = "Muestras_CVMVF"
FC_PATH = os.path.join(GDB, FC_NAME)
LAYER_NAME = "Muestras CVMVF"


def rgb(r, g, b, a=100):
    c = arcpy.cim.CreateCIMObjectFromClassName("CIMRGBColor", "V3")
    c.values = [float(r), float(g), float(b), float(a)]
    return c


def set_red_point_symbol(lyr, size_pt=20.0):
    """Circulo rojo vectorial (CIMVectorMarker) — se dibuja de forma fiable."""
    import math

    fill = arcpy.cim.CreateCIMObjectFromClassName("CIMSolidFill", "V3")
    fill.color = rgb(255, 0, 0, 100)
    stroke = arcpy.cim.CreateCIMObjectFromClassName("CIMSolidStroke", "V3")
    stroke.color = rgb(80, 0, 0, 100)
    stroke.width = 1.5
    stroke.enable = True
    poly_sym = arcpy.cim.CreateCIMObjectFromClassName("CIMPolygonSymbol", "V3")
    poly_sym.symbolLayers = [stroke, fill]

    pts = [
        arcpy.Point(math.cos(2 * math.pi * i / 32.0), math.sin(2 * math.pi * i / 32.0))
        for i in range(33)
    ]
    circ = arcpy.Polygon(arcpy.Array(pts))
    graphic = arcpy.cim.CreateCIMObjectFromClassName("CIMMarkerGraphic", "V3")
    # Debe ser el simbolo directo: un CIMSymbolReference aqui se serializa como null
    graphic.symbol = poly_sym
    graphic.geometry = circ

    vm = arcpy.cim.CreateCIMObjectFromClassName("CIMVectorMarker", "V3")
    vm.size = float(size_pt)
    vm.enable = True
    vm.frame = arcpy.Extent(-1, -1, 1, 1)
    vm.markerGraphics = [graphic]

    pt = arcpy.cim.CreateCIMObjectFromClassName("CIMPointSymbol", "V3")
    pt.symbolLayers = [vm]
    renderer = arcpy.cim.CreateCIMObjectFromClassName("CIMSimpleRenderer", "V3")
    symref = arcpy.cim.CreateCIMObjectFromClassName("CIMSymbolReference", "V3")
    symref.symbol = pt
    renderer.symbol = symref
    cim = lyr.getDefinition("V3")
    cim.renderer = renderer
    cim.minScale = 0
    cim.maxScale = 0
    lyr.setDefinition(cim)


def set_red_point_symbol_simple(lyr, size_pt=9.0):
    try:
        sym = lyr.symbology
        if hasattr(sym, "renderer") and sym.renderer.type == "SimpleRenderer":
            sym.renderer.symbol.color = {"RGB": [220, 0, 0, 100]}
            sym.renderer.symbol.size = size_pt
            lyr.symbology = sym
            return True
    except Exception as ex:
        print(f"  Aviso symbology API: {ex}")
    return False


def configure_labels_codigo(lyr):
    try:
        fields = [f.name for f in arcpy.ListFields(lyr.dataSource)]
        if "Codigo" not in fields:
            return
        lyr.showLabels = True
        cim = lyr.getDefinition("V3")
        if not cim.labelClasses:
            return
        lc = cim.labelClasses[0]
        lc.visibility = True
        # Python [Campo] es mas fiable que Arcade en algunos layouts
        lc.expressionEngine = "Python"
        lc.expression = "[Codigo]"
        try:
            lc.minScale = 0
            lc.maxScale = 0
        except Exception:
            pass
        ts = lc.textSymbol.symbol
        ts.height = 12
        ts.fontFamilyName = "Arial"
        ts.fontStyleName = "Bold"
        for sl in getattr(ts, "symbolLayers", []) or []:
            if type(sl).__name__ == "CIMSolidFill":
                sl.color = rgb(0, 0, 0, 100)
        fill = arcpy.cim.CreateCIMObjectFromClassName("CIMSolidFill", "V3")
        fill.color = rgb(255, 255, 255, 100)
        poly = arcpy.cim.CreateCIMObjectFromClassName("CIMPolygonSymbol", "V3")
        poly.symbolLayers = [fill]
        try:
            ts.haloSize = 2.0
            ts.haloSymbol = poly
        except Exception:
            pass
        lyr.setDefinition(cim)
        lyr.showLabels = True
        print("  OK Etiquetas: [Codigo]")
    except Exception as ex:
        print(f"  Aviso etiquetas: {ex}")


def import_excel_to_fc():
    """Lee el Excel y crea puntos EPSG:9377 en la GDB."""
    if not os.path.exists(EXCEL):
        raise FileNotFoundError(EXCEL)
    if not arcpy.Exists(GDB):
        raise FileNotFoundError(GDB)

    import openpyxl

    arcpy.env.overwriteOutput = True

    wb = openpyxl.load_workbook(EXCEL, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise RuntimeError("Excel vacio")

    header = [str(h).strip() if h is not None else f"COL{i}" for i, h in enumerate(rows[0])]

    def col(*cands):
        lower = {h.lower(): i for i, h in enumerate(header)}
        for c in cands:
            if c.lower() in lower:
                return lower[c.lower()]
        for h, i in lower.items():
            for c in cands:
                if c.lower() in h:
                    return i
        return None

    i_lat = col("Latitud", "lat", "latitude")
    i_lon = col("Longitud", "lon", "longitude", "long")
    i_elev = col("Elevacion", "Elevación", "elev", "elevation")
    i_cod = col("Codigo", "Código", "code")
    if i_lat is None or i_lon is None:
        raise RuntimeError(f"No hay Latitud/Longitud en {header}")

    records = []
    for row in rows[1:]:
        if row is None:
            continue
        try:
            lat = float(row[i_lat]) if row[i_lat] is not None else None
            lon = float(row[i_lon]) if row[i_lon] is not None else None
        except (TypeError, ValueError):
            continue
        if lat is None or lon is None:
            continue
        elev = None
        if i_elev is not None and row[i_elev] is not None:
            try:
                elev = float(row[i_elev])
            except (TypeError, ValueError):
                elev = None
        codigo = row[i_cod] if i_cod is not None else None
        attrs = {}
        for i, h in enumerate(header):
            if i in (i_lat, i_lon):
                continue
            val = row[i] if i < len(row) else None
            if val is None:
                attrs[h] = None
            elif hasattr(val, "isoformat"):
                attrs[h] = val.isoformat()[:50]
            else:
                attrs[h] = str(val)[:250]
        records.append((lon, lat, elev, codigo, attrs))

    print(f"  Filas con coordenadas: {len(records)}")
    if not records:
        raise RuntimeError("Ningun punto con Latitud/Longitud validas")

    def pick(attrs, *keys):
        lower = {k.lower(): v for k, v in attrs.items()}
        for k in keys:
            for lk, v in lower.items():
                if k.lower() in lk:
                    return v
        return None

    sr_wgs = arcpy.SpatialReference(4326)
    sr_out = arcpy.SpatialReference(9377)
    tmp_name = "Muestras_CVMVF_wgs_tmp"
    tmp_wgs = os.path.join(GDB, tmp_name)
    if arcpy.Exists(tmp_wgs):
        arcpy.management.Delete(tmp_wgs)
    if arcpy.Exists(FC_PATH):
        arcpy.management.Delete(FC_PATH)

    arcpy.management.CreateFeatureclass(
        GDB, tmp_name, "POINT", spatial_reference=sr_wgs, has_z="ENABLED"
    )
    arcpy.management.AddField(tmp_wgs, "Codigo", "TEXT", field_length=50)
    arcpy.management.AddField(tmp_wgs, "Elevacion", "DOUBLE")
    arcpy.management.AddField(tmp_wgs, "Tipo_Roca", "TEXT", field_length=100)
    arcpy.management.AddField(tmp_wgs, "Clasificacion", "TEXT", field_length=100)
    arcpy.management.AddField(tmp_wgs, "Localidad", "TEXT", field_length=250)
    arcpy.management.AddField(tmp_wgs, "Descripcion", "TEXT", field_length=250)

    with arcpy.da.InsertCursor(
        tmp_wgs,
        [
            "SHAPE@XY",
            "SHAPE@Z",
            "Codigo",
            "Elevacion",
            "Tipo_Roca",
            "Clasificacion",
            "Localidad",
            "Descripcion",
        ],
    ) as cur:
        for lon, lat, elev, codigo, attrs in records:
            z = elev if elev is not None else 0.0
            cur.insertRow(
                [
                    (lon, lat),
                    z,
                    str(codigo)[:50] if codigo else None,
                    elev,
                    (pick(attrs, "Tipo de Roca", "Tipo_de_Roca") or "" )[:100] or None,
                    (pick(attrs, "Clasificacion", "Clasificación") or "")[:100] or None,
                    (pick(attrs, "Localidad") or "")[:250] or None,
                    (pick(attrs, "Descripcion", "Descripción") or "")[:250] or None,
                ]
            )

    print("  Project -> EPSG:9377...")
    arcpy.management.Project(tmp_wgs, FC_PATH, sr_out)
    n = int(arcpy.management.GetCount(FC_PATH)[0])
    print(f"  OK {FC_PATH} ({n} features)")
    try:
        arcpy.management.Delete(tmp_wgs)
    except Exception:
        pass
    return FC_PATH, n


def add_to_maps(aprx, fc_path):
    targets = []
    for name in ("Map_Estructural", "Map"):
        maps = aprx.listMaps(name)
        if maps:
            targets.append(maps[0])
            break
    for m in aprx.listMaps("Mapa_Litologico"):
        targets.append(m)

    if not targets:
        raise RuntimeError("No se encontro Map_Estructural ni Map")

    added = []
    for amap in targets:
        for lyr in list(amap.listLayers()):
            if lyr.name == LAYER_NAME or lyr.name == FC_NAME:
                try:
                    amap.removeLayer(lyr)
                except Exception:
                    pass
        lyr = amap.addDataFromPath(fc_path)
        try:
            lyr.name = LAYER_NAME
        except Exception:
            pass
        lyr.visible = True
        try:
            set_red_point_symbol(lyr, 9.0)
            print(f"  OK CIM rojo en {amap.name}")
        except Exception as ex:
            print(f"  Aviso CIM rojo: {ex}")
            if set_red_point_symbol_simple(lyr, 9.0):
                print(f"  OK symbology API rojo en {amap.name}")
        configure_labels_codigo(lyr)
        try:
            ref = amap.listLayers()[0]
            if lyr != ref:
                amap.moveLayer(ref, lyr, "BEFORE")
        except Exception:
            pass
        added.append(amap.name)
        print(f"  OK Capa '{LAYER_NAME}' en {amap.name}")
    return added


def main():
    print("=" * 60)
    print("Importar Muestras CVMVF -> puntos rojos")
    print("=" * 60)

    print("\n[1] Excel -> feature class")
    fc, n = import_excel_to_fc()

    print("\n[2] Anadir a mapas")
    aprx = arcpy.mp.ArcGISProject(APRX_PATH)
    maps = add_to_maps(aprx, fc)

    print("\n[3] Guardar APRX")
    aprx.save()
    print(f"  OK {APRX_PATH}")
    print(f"\nListo: {n} puntos en {', '.join(maps)}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        traceback.print_exc()
        sys.exit(1)
