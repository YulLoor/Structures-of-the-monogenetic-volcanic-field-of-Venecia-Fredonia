# -*- coding: utf-8 -*-
"""
Importa 01_Insumos/Datos_Geologia_Estructural_RHR.xlsx como puntos
(WGS84 X/Y -> EPSG:9377) y los dibuja en Map_Estructural con el
simbolo T de rumbo-buzamiento (convencion RHR: tic a la derecha del rumbo).

Ejecutar:
  "C:\\Program Files\\ArcGIS\\Pro\\bin\\Python\\envs\\arcgispro-py3\\python.exe" ^
    04_Scripts\\importar_rumbo_buzamiento.py
"""

from __future__ import annotations

import math
import os
import sys
import traceback

import arcpy

ROOT = r"C:\PROYECTO_GIS_VF_Antigraviti"
APRX_PATH = os.path.join(ROOT, "Venecia_Fredonia_Analisis_Estructural.aprx")
GDB = os.path.join(
    ROOT,
    "Venecia_Fredonia_Analisis_Estructural",
    "Venecia_Fredonia_Analisis_Estructural.gdb",
)
EXCEL = os.path.join(ROOT, "01_Insumos", "Datos_Geologia_Estructural_RHR.xlsx")
FC_NAME = "Datos_Estructural_RHR"
FC_PATH = os.path.join(GDB, FC_NAME)
LAYER_NAME = "Rumbo y buzamiento"
SYMBOL_SIZE_PT = 16.0
DIP_LABEL_PT = 8.0


def rgb(r, g, b, a=100):
    c = arcpy.cim.CreateCIMObjectFromClassName("CIMRGBColor", "V3")
    c.values = [float(r), float(g), float(b), float(a)]
    return c


def _poly_symbol():
    fill = arcpy.cim.CreateCIMObjectFromClassName("CIMSolidFill", "V3")
    fill.color = rgb(0, 0, 0, 100)
    fill.enable = True
    stroke = arcpy.cim.CreateCIMObjectFromClassName("CIMSolidStroke", "V3")
    stroke.color = rgb(0, 0, 0, 100)
    stroke.width = 0.4
    stroke.enable = True
    poly = arcpy.cim.CreateCIMObjectFromClassName("CIMPolygonSymbol", "V3")
    poly.symbolLayers = [stroke, fill]
    return poly


def _rect(xmin, ymin, xmax, ymax):
    pts = [
        arcpy.Point(xmin, ymin),
        arcpy.Point(xmin, ymax),
        arcpy.Point(xmax, ymax),
        arcpy.Point(xmax, ymin),
        arcpy.Point(xmin, ymin),
    ]
    return arcpy.Polygon(arcpy.Array(pts))


def _graphic_poly(geom):
    graphic = arcpy.cim.CreateCIMObjectFromClassName("CIMMarkerGraphic", "V3")
    graphic.symbol = _poly_symbol()
    graphic.geometry = geom
    return graphic


def strike_dip_t_symbol(size_pt=SYMBOL_SIZE_PT):
    """
    T geologico en poligonos (se dibuja de forma fiable): rumbo vertical
    (N-S a rotacion 0) y tic de buzamiento hacia el este (derecha).
    Con rotacion geografica = Rumbo, el tic queda a la derecha (RHR).
    """
    w = 0.11
    strike = _graphic_poly(_rect(-w, -1.0, w, 1.0))
    tick = _graphic_poly(_rect(0.0, -w, 0.78, w))
    vm = arcpy.cim.CreateCIMObjectFromClassName("CIMVectorMarker", "V3")
    vm.size = float(size_pt)
    vm.enable = True
    vm.frame = arcpy.Extent(-1, -1, 1, 1)
    vm.markerGraphics = [strike, tick]
    try:
        vm.anchorPoint = arcpy.Point(0, 0)
    except Exception:
        pass
    pt = arcpy.cim.CreateCIMObjectFromClassName("CIMPointSymbol", "V3")
    pt.symbolLayers = [vm]
    try:
        pt.angleAlignment = "Map"
    except Exception:
        pass
    ref = arcpy.cim.CreateCIMObjectFromClassName("CIMSymbolReference", "V3")
    ref.symbol = pt
    return ref


def apply_rotation_variable(renderer, field="Rotacion"):
    rot = arcpy.cim.CreateCIMObjectFromClassName("CIMRotationVisualVariable", "V3")
    try:
        rot.rotationTypeZ = "Geographic"
    except Exception:
        pass
    expr = arcpy.cim.CreateCIMObjectFromClassName("CIMExpressionInfo", "V3")
    expr.expression = f"$feature.{field}"
    expr.title = field
    try:
        expr.returnType = "Numeric"
    except Exception:
        pass
    info = arcpy.cim.CreateCIMObjectFromClassName("CIMVisualVariableInfo", "V3")
    info.visualVariableInfoType = "Expression"
    info.expression = f"$feature.{field}"
    info.valueExpressionInfo = expr
    try:
        rot.visualVariableInfoZ = info
    except Exception as ex:
        print(f"  Aviso visualVariableInfoZ: {ex}")
        rot.visualVariableInfo = info
    existing = [
        v
        for v in (getattr(renderer, "visualVariables", None) or [])
        if type(v).__name__ != "CIMRotationVisualVariable"
    ]
    renderer.visualVariables = existing + [rot]


def set_strike_dip_symbol(lyr):
    rend = arcpy.cim.CreateCIMObjectFromClassName("CIMSimpleRenderer", "V3")
    rend.symbol = strike_dip_t_symbol()
    rend.label = LAYER_NAME
    rend.description = LAYER_NAME
    apply_rotation_variable(rend, "Rotacion")
    cim = lyr.getDefinition("V3")
    cim.renderer = rend
    cim.minScale = 0
    cim.maxScale = 0
    lyr.setDefinition(cim)
    lyr.visible = True
    print(f"  OK Simbolo T rumbo-buzamiento {SYMBOL_SIZE_PT:.0f} pt")


def configure_dip_labels(lyr):
    cim = lyr.getDefinition("V3")
    if not cim.labelClasses:
        lc = arcpy.cim.CreateCIMObjectFromClassName("CIMLabelClass", "V3")
        cim.labelClasses = [lc]
    lc = cim.labelClasses[0]
    lc.visibility = True
    lc.name = "Buzamiento"
    lc.expressionEngine = "Python"
    lc.expression = "[Buzamiento]"
    try:
        lc.minScale = 0
        lc.maxScale = 0
    except Exception:
        pass
    try:
        lc.minimumScale = 0.0
        lc.maximumScale = 0.0
    except Exception:
        pass
    ts = lc.textSymbol.symbol
    ts.height = DIP_LABEL_PT
    ts.fontFamilyName = "Arial"
    ts.fontStyleName = "Bold"
    for sl in getattr(ts, "symbolLayers", []) or []:
        if type(sl).__name__ == "CIMSolidFill":
            sl.color = rgb(0, 0, 0, 100)
    halo_fill = arcpy.cim.CreateCIMObjectFromClassName("CIMSolidFill", "V3")
    halo_fill.color = rgb(255, 255, 255, 100)
    halo = arcpy.cim.CreateCIMObjectFromClassName("CIMPolygonSymbol", "V3")
    halo.symbolLayers = [halo_fill]
    try:
        ts.haloSize = 1.2
        ts.haloSymbol = halo
    except Exception:
        pass
    try:
        mp = lc.maplexLabelPlacementProperties
        mp.featureType = "Point"
        mp.pointPlacementMethod = "AroundPoint"
        mp.primaryOffset = 6.0
        mp.primaryOffsetUnit = "Point"
        mp.thinDuplicateLabels = False
        mp.neverRemoveLabel = True
    except Exception as ex:
        print(f"  Aviso Maplex: {ex}")
    lyr.setDefinition(cim)
    lyr.showLabels = True
    print(f"  OK Etiquetas [Buzamiento] {DIP_LABEL_PT:.0f} pt")


def _col(header, *cands):
    lower = {h.lower(): i for i, h in enumerate(header)}
    for c in cands:
        if c.lower() in lower:
            return lower[c.lower()]
    for h, i in lower.items():
        for c in cands:
            if c.lower() in h:
                return i
    return None


def import_excel_to_fc():
    if not os.path.exists(EXCEL):
        raise FileNotFoundError(EXCEL)
    if not arcpy.Exists(GDB):
        raise FileNotFoundError(GDB)

    import openpyxl

    arcpy.env.overwriteOutput = True
    wb = openpyxl.load_workbook(EXCEL, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise RuntimeError("Excel vacio")

    header = [str(h).strip() if h is not None else f"COL{i}" for i, h in enumerate(rows[0])]
    i_x = _col(header, "X", "Longitud", "lon", "longitude")
    i_y = _col(header, "Y", "Latitud", "lat", "latitude")
    i_id = _col(header, "ID")
    i_rumbo = _col(header, "Rumbo")
    i_buz = _col(header, "Buzamiento")
    i_db = _col(header, "DB")
    i_est = _col(header, "Estación", "Estacion")
    i_elev = _col(header, "Elevación", "Elevacion")
    i_ubi = _col(header, "Ubicación", "Ubicacion")
    i_tipo = _col(header, "Tipo de Estructura", "Tipo")
    i_fam = _col(header, "Familia")
    if i_x is None or i_y is None or i_rumbo is None or i_buz is None:
        raise RuntimeError(f"Faltan X/Y/Rumbo/Buzamiento en {header}")

    records = []
    for row in rows[1:]:
        if row is None:
            continue
        try:
            lon = float(row[i_x])
            lat = float(row[i_y])
            rumbo = float(row[i_rumbo])
            buz = float(row[i_buz])
        except (TypeError, ValueError):
            continue
        rec_id = row[i_id] if i_id is not None else len(records) + 1
        try:
            rec_id = int(rec_id)
        except (TypeError, ValueError):
            rec_id = len(records) + 1
        elev = None
        if i_elev is not None and row[i_elev] is not None:
            try:
                elev = float(row[i_elev])
            except (TypeError, ValueError):
                elev = None
        records.append(
            {
                "id": rec_id,
                "lon": lon,
                "lat": lat,
                "rumbo": rumbo,
                "buz": buz,
                "db": (str(row[i_db]).strip() if i_db is not None and row[i_db] else None),
                "est": (str(row[i_est]).strip() if i_est is not None and row[i_est] else None),
                "elev": elev,
                "ubi": (str(row[i_ubi]).strip() if i_ubi is not None and row[i_ubi] else None),
                "tipo": (str(row[i_tipo]).strip() if i_tipo is not None and row[i_tipo] else None),
                "fam": (str(row[i_fam]).strip() if i_fam is not None and row[i_fam] else None),
            }
        )

    print(f"  Filas con rumbo-buzamiento: {len(records)}")
    if not records:
        raise RuntimeError("Ningun registro valido")

    # Offset en metros para puntos coincidentes (Morro Alegre 1 y 2)
    seen_xy = {}
    for rec in records:
        key = (round(rec["lon"], 6), round(rec["lat"], 6))
        seen_xy[key] = seen_xy.get(key, 0) + 1
        rec["_dup"] = seen_xy[key] - 1

    sr_wgs = arcpy.SpatialReference(4326)
    sr_out = arcpy.SpatialReference(9377)
    tmp_name = "RHR_wgs_tmp"
    tmp_wgs = os.path.join(GDB, tmp_name)
    for p in (tmp_wgs, FC_PATH):
        if arcpy.Exists(p):
            try:
                arcpy.management.Delete(p)
            except Exception as ex:
                print(f"  Aviso delete {p}: {ex}")

    arcpy.management.CreateFeatureclass(
        GDB, tmp_name, "POINT", spatial_reference=sr_wgs
    )
    fields = [
        ("ID_med", "SHORT"),
        ("Rumbo", "DOUBLE"),
        ("Buzamiento", "DOUBLE"),
        ("DB", "TEXT", 8),
        ("Rotacion", "DOUBLE"),
        ("Estacion", "TEXT", 80),
        ("Elevacion", "DOUBLE"),
        ("Ubicacion", "TEXT", 120),
        ("Tipo_Estr", "TEXT", 80),
        ("Familia", "TEXT", 40),
    ]
    for spec in fields:
        if len(spec) == 2:
            arcpy.management.AddField(tmp_wgs, spec[0], spec[1])
        else:
            arcpy.management.AddField(tmp_wgs, spec[0], spec[1], field_length=spec[2])

    insert_fields = [
        "SHAPE@XY",
        "ID_med",
        "Rumbo",
        "Buzamiento",
        "DB",
        "Rotacion",
        "Estacion",
        "Elevacion",
        "Ubicacion",
        "Tipo_Estr",
        "Familia",
    ]
    with arcpy.da.InsertCursor(tmp_wgs, insert_fields) as cur:
        for rec in records:
            lon, lat = rec["lon"], rec["lat"]
            if rec["_dup"]:
                # ~12 m por duplicado, en circulo, para que los T no se tapen
                ddeg = 12.0 / 111320.0
                ang = rec["_dup"] * (2.0 * math.pi / 3.0)
                lon = lon + ddeg * math.cos(ang) / math.cos(math.radians(lat))
                lat = lat + ddeg * math.sin(ang)
            cur.insertRow(
                [
                    (lon, lat),
                    rec["id"],
                    rec["rumbo"],
                    rec["buz"],
                    rec["db"],
                    rec["rumbo"],
                    rec["est"],
                    rec["elev"],
                    rec["ubi"],
                    rec["tipo"],
                    rec["fam"],
                ]
            )

    print("  Project -> EPSG:9377...")
    arcpy.management.Project(tmp_wgs, FC_PATH, sr_out)
    n = int(arcpy.management.GetCount(FC_PATH)[0])
    print(f"  OK {FC_PATH} ({n} features)")
    try:
        arcpy.management.Delete(tmp_wgs)
    except Exception:
        pass
    return FC_PATH, n


def add_to_map(aprx, fc_path):
    amap = None
    for name in ("Map_Estructural", "Map"):
        maps = aprx.listMaps(name)
        if maps:
            amap = maps[0]
            break
    if amap is None:
        raise RuntimeError("No se encontro Map_Estructural")

    for lyr in list(amap.listLayers()):
        if lyr.name == LAYER_NAME or lyr.name == FC_NAME:
            try:
                amap.removeLayer(lyr)
            except Exception:
                pass
    lyr = amap.addDataFromPath(fc_path)
    try:
        lyr.name = LAYER_NAME
    except Exception:
        pass
    lyr.visible = True
    set_strike_dip_symbol(lyr)
    cim_chk = lyr.getDefinition("V3")
    vv = (cim_chk.renderer.visualVariables or [None])[0]
    if vv is not None:
        info = getattr(vv, "visualVariableInfoZ", None)
        print(
            "  Rotacion Z:",
            getattr(vv, "rotationTypeZ", None),
            getattr(info, "visualVariableInfoType", None),
            getattr(info, "expression", None),
        )
    configure_dip_labels(lyr)
    try:
        muestras = None
        for L in amap.listLayers():
            if L.name == "Muestras CVMVF":
                muestras = L
                break
        if muestras is not None:
            amap.moveLayer(muestras, lyr, "BEFORE")
        else:
            top = amap.listLayers()[0]
            if lyr != top:
                amap.moveLayer(top, lyr, "BEFORE")
    except Exception as ex:
        print(f"  Aviso orden capas: {ex}")
    print(f"  OK Capa {LAYER_NAME} en {amap.name}")
    return amap.name


def main():
    print("=" * 60)
    print("Importar rumbo-buzamiento (Datos_Geologia_Estructural_RHR)")
    print("=" * 60)
    if not os.path.exists(APRX_PATH):
        raise FileNotFoundError(APRX_PATH)
    print("\n[1] Excel -> GDB")
    fc, n = import_excel_to_fc()
    print(f"\n[2] Anadir a Map_Estructural ({n} medidas)")
    aprx = arcpy.mp.ArcGISProject(APRX_PATH)
    add_to_map(aprx, fc)
    print("\n[3] Guardar APRX")
    try:
        aprx.save()
        print(f"  OK {APRX_PATH}")
    except Exception as ex:
        print(f"  ERROR al guardar: {ex}")
        print("  Cierra ArcGIS Pro SIN guardar y vuelve a ejecutar el script.")
        raise
    print("\nListo. Si Pro tenia el proyecto abierto, cierra SIN guardar y reabre.")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        traceback.print_exc()
        sys.exit(1)
